from __future__ import annotations

import csv
import copy
import base64
import html
import io
import json
import platform
import subprocess
import sys
import traceback
import tempfile
import threading
import zipfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import webview
from Bio import SeqIO
from Bio.Seq import Seq
from openpyxl import load_workbook
from webview.dom import DOMEventHandler
from webview.menu import Menu, MenuAction, MenuSeparator

from plasmid_verify.assembly import (
    TYPE_IIS_ENZYMES,
    AssemblyError,
    assemble_gibson,
    assemble_golden_gate,
)
from plasmid_verify.fasta import clean_sequence, infer_target_id, parse_sequence_file, to_fasta
from plasmid_verify.golden_gate_design import plan_from_annotated_snapgene, run_annotated_snapgene_design, run_plan
from plasmid_verify.models import AssemblyResult, Mutation, VerificationResult
from plasmid_verify.verify import verify_consensus
from mac_app import local_updates, github_updates, design_library


APP_NAME = "CloningCompanion"
APP_VERSION = "1.8.2"
PROJECT_EXTENSION = "plasmidverify"
SEQUENCE_TYPES = ("Sequence files (*.fasta;*.fa;*.fna;*.fas;*.dna)", "All files (*.*)")
INSERT_TYPES = (
    "Insert sequences or tables (*.fasta;*.fa;*.fna;*.fas;*.dna;*.csv;*.tsv;*.xlsx;*.txt)",
    "All files (*.*)",
)
PROJECT_TYPES = ("CloningCompanion projects (*.plasmidverify)",)
DESIGN_MAP_TYPES = ("Annotated SnapGene maps (*.dna)", "All files (*.*)")


def _project_name(value: Optional[str]) -> str:
    """Keep explicit names; timestamp unnamed projects in the Mac's local time."""
    return value if value and value.strip() else datetime.now().astimezone().strftime("%Y-%m-%d %H-%M-%S")


def _same_design_inputs(first: Dict[str, Any], second: Dict[str, Any]) -> bool:
    a, b = first.get("goldenGateDesign") or {}, second.get("goldenGateDesign") or {}
    return (a.get("enzyme") == b.get("enzyme") and a.get("variableTexts", {}) == b.get("variableTexts", {}) and
            (a.get("map") or {}).get("embeddedSnapGene") == (b.get("map") or {}).get("embeddedSnapGene"))


def resource_path(relative: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
    return base / relative


def application_support() -> Path:
    directory = Path.home() / "Library" / "Application Support" / APP_NAME
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _map_overview(path: Path, title: str) -> str:
    """A whole-sequence overview accompanying the editable SnapGene file."""
    record = SeqIO.read(path, "snapgene")
    cores = []
    for feature in record.features:
        label = str(feature.qualifiers.get("label", [""])[0])
        if label.startswith(("[", "{")):
            cores.append((label, feature))
    height = 170 + len(cores) * 30
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="1100" height="{height}" viewBox="0 0 1100 {height}">',
           '<rect width="100%" height="100%" fill="white"/>',
           f'<text x="40" y="35" font-family="Arial,sans-serif" font-size="20" fill="#18382c">{html.escape(title)}</text>',
           f'<text x="40" y="59" font-family="Arial,sans-serif" font-size="13" fill="#627369">{len(record):,} bp · circular sequence shown linearly</text>',
           '<line x1="40" y1="105" x2="1060" y2="105" stroke="#bdc8c1" stroke-width="2"/>']
    for index, (label, feature) in enumerate(cores):
        color = ("#478cba", "#559d79", "#9875bc", "#cf9254")[index % 4]
        for part in feature.location.parts:
            x = 40 + int(part.start) / len(record) * 1020
            width = max(1, len(part) / len(record) * 1020)
            svg.append(f'<rect x="{x:.2f}" y="89" width="{width:.2f}" height="32" rx="3" fill="{color}"/>')
        y = 174 + index * 30
        svg.extend([f'<rect x="40" y="{y - 13}" width="13" height="13" rx="2" fill="{color}"/>',
                    f'<text x="65" y="{y}" font-family="Arial,sans-serif" font-size="14" fill="#243c30">{html.escape(label)} · {len(feature):,} bp</text>'])
    svg.extend([f'<text x="40" y="142" font-family="Arial,sans-serif" font-size="12" fill="#627369">1</text>',
                f'<text x="1060" y="142" text-anchor="end" font-family="Arial,sans-serif" font-size="12" fill="#627369">{len(record):,}</text>', '</svg>'])
    return "".join(svg)


def _design_previews(result: Dict[str, Any]) -> Dict[str, Any]:
    """Embed local graphics so the native WebKit view needs no file permissions."""
    graphics = []
    def add(path, title, group, mime="image/png", source_path=None):
        path = Path(path)
        graphics.append({"title": title, "group": group, "path": str(path), "sourcePath": str(source_path or path),
                         "src": f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode("ascii")})
    junctions = {row["junction_id"]: row["selected_fusion"] for row in result.get("junctions", [])}
    if result.get("interactiveMap"):
        add(result["interactiveMap"]["path"], "Linear map", "Linear map")
        map_graphic = graphics[-1]
        map_graphic["junctions"] = []
        for junction in result["interactiveMap"]["junctions"]:
            detail_index = len(graphics)
            add(junction["detailPath"], "Junction " + junction["label"], "Junction details")
            map_graphic["junctions"].append({key: value for key, value in junction.items() if key != "detailPath"})
            map_graphic["junctions"][-1]["graphicIndex"] = detail_index
    elif result.get("plasmidJunctionFigure"):
        add(result["plasmidJunctionFigure"], "Linear map", "Linear map")
    for path in result.get("junctionScreencaps", []):
        path = Path(path)
        title = f"{path.stem.replace('_', ' ').capitalize()} · {junctions.get(path.stem, '')}"
        add(path, title, "Fragment ends")
    result["graphics"] = graphics
    return result


def _first_sequence(path: str) -> Tuple[str, str]:
    data = Path(path).read_bytes()
    record = parse_sequence_file(data, Path(path).name)[0]
    return record.name, record.sequence


def _insert_rows_from_values(
    fieldnames: Sequence[Any], values: Iterable[Sequence[Any]], source: str
) -> List[Dict[str, Any]]:
    """Parse normalized rows with target_id and sequence/DNA sequence columns."""
    if not fieldnames:
        raise ValueError("Insert table needs a header row")
    headers = [str(header or "").strip() for header in fieldnames]
    columns = {header.lower().replace(" ", "_").replace("-", "_"): index for index, header in enumerate(headers) if header}
    target_column = next((columns[key] for key in ("target_id", "target", "clone_id", "clone") if key in columns), None)
    sequence_column = next((columns[key] for key in ("sequence", "dna_sequence", "dna", "insert_sequence") if key in columns), None)
    if target_column is None or sequence_column is None:
        raise ValueError("Insert table requires target_id (or clone_id) and sequence (or dna_sequence) columns")
    name_column = next((columns[key] for key in ("name", "insert_name", "part", "id") if key in columns), None)
    order_column = next((columns[key] for key in ("order", "assembly_order", "insert_order") if key in columns), None)
    rows: List[Dict[str, Any]] = []
    for row_number, values_row in enumerate(values, 2):
        values_list = list(values_row)
        def value_at(column: Optional[int]) -> str:
            return str(values_list[column] or "").strip() if column is not None and column < len(values_list) else ""
        target_id = value_at(target_column)
        raw_sequence = value_at(sequence_column)
        if not target_id and not raw_sequence:
            continue
        if not target_id or not raw_sequence:
            raise ValueError(f"{source}, row {row_number}: target ID and sequence are both required")
        sequence = clean_sequence(raw_sequence)
        label = value_at(name_column) or f"insert_{len(rows) + 1}"
        raw_order = value_at(order_column)
        try:
            order = int(raw_order) if raw_order else 0
        except ValueError as exc:
            raise ValueError(f"{source}, row {row_number}: order must be a whole number") from exc
        rows.append({
            "path": "", "filename": f"{source} · {label}", "recordName": label,
            "length": len(sequence), "targetId": target_id, "sequence": sequence,
            "order": order, "noncoding5": 0, "noncoding3": 0,
        })
    if not rows:
        raise ValueError("Insert table contains no insert rows")
    return rows


def _table_insert_rows(text: str, source: str) -> List[Dict[str, Any]]:
    """Parse a CSV/TSV insert table with target_id and sequence columns."""
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel_tab if "\t" in text.partition("\n")[0] else csv.excel
    reader = csv.reader(text.splitlines(), dialect=dialect)
    return _insert_rows_from_values(next(reader, []), reader, source)


def _xlsx_insert_rows(path: str) -> List[Dict[str, Any]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = next(values, None)
        return _insert_rows_from_values(headers or [], values, Path(path).name)
    except Exception as exc:
        if isinstance(exc, ValueError):
            raise
        raise ValueError(f"Could not read Excel insert table: {exc}") from exc
    finally:
        try:
            workbook.close()
        except UnboundLocalError:
            pass


def _safe_filename(value: str) -> str:
    cleaned = "".join(character if character.isalnum() or character in "-_." else "_" for character in value)
    return cleaned.strip("._") or "target"


def _alignment_window(result: VerificationResult, mutation: Mutation, flank: int = 24) -> Dict[str, str]:
    target_seen = 0
    center = 0
    for index, base in enumerate(result.aligned_target):
        if base != "-":
            target_seen += 1
        if target_seen >= mutation.target_position:
            center = index
            break
    start = max(0, center - flank)
    end = min(len(result.aligned_target), center + flank + 1)
    expected = result.aligned_target[start:end]
    observed = result.aligned_observed[start:end]
    return {
        "target": expected,
        "marks": "".join("|" if a == b else "•" for a, b in zip(expected, observed)),
        "observed": observed,
    }


def _codon_zoom(assembly: AssemblyResult, result: VerificationResult, mutation: Mutation) -> Optional[Dict[str, Any]]:
    """Return the affected coding codon in both target and WPS coordinates."""
    target_index = mutation.target_position - 1
    if target_index < 0 or target_index >= len(assembly.annotations):
        return None
    annotation = assembly.annotations[target_index]
    if annotation.region != "coding" or annotation.feature_start is None:
        return None
    codon_start = target_index - ((target_index - annotation.feature_start) % 3)
    expected = assembly.sequence[codon_start : codon_start + 3]
    if len(expected) != 3:
        return None
    target_to_observed: Dict[int, str] = {}
    position = 0
    for target_base, observed_base in zip(result.aligned_target, result.aligned_observed):
        if target_base != "-":
            target_to_observed[position] = observed_base
            position += 1
    observed = "".join(target_to_observed.get(index, "-") for index in range(codon_start, codon_start + 3))
    expected_aa = str(Seq(expected).translate())
    observed_aa = "?" if "-" in observed or "N" in observed else str(Seq(observed).translate())
    return {
        "expected": expected,
        "observed": observed,
        "expectedAa": expected_aa,
        "observedAa": observed_aa,
        "aaPosition": ((target_index - annotation.feature_start) // 3) + 1,
        "codonStart": codon_start + 1,
    }


def _annotation_runs(assembly: AssemblyResult) -> List[Dict[str, Any]]:
    if not assembly.annotations:
        return []
    result: List[Dict[str, Any]] = []
    start = 0
    current = assembly.annotations[0]
    for index, annotation in enumerate(assembly.annotations[1:], 1):
        if (annotation.region, annotation.source) != (current.region, current.source):
            result.append({"start": start, "end": index, "region": current.region, "source": current.source})
            start = index
            current = annotation
    result.append(
        {"start": start, "end": len(assembly.annotations), "region": current.region, "source": current.source}
    )
    return result


def _serialize_result(assembly: AssemblyResult, verification: VerificationResult) -> Dict[str, Any]:
    mutations = []
    for mutation in verification.mutations:
        item = mutation.as_dict()
        item["alignment"] = _alignment_window(verification, mutation)
        item["codon"] = _codon_zoom(assembly, verification, mutation)
        mutations.append(item)
    return {
        "targetId": verification.target_id,
        "verdict": verification.verdict,
        "method": assembly.method,
        "targetLength": verification.expected_length,
        "observedLength": verification.observed_length,
        "orientation": verification.orientation,
        "rotation": verification.rotation,
        "counts": verification.counts,
        "mutations": mutations,
        "notes": assembly.notes,
        "runs": _annotation_runs(assembly),
        "sequence": assembly.sequence,
    }


class NativeAPI:
    def __init__(self) -> None:
        self.window: Optional[webview.Window] = None
        self.latest_results: Dict[str, Tuple[AssemblyResult, VerificationResult]] = {}
        self.current_project: Optional[str] = None
        self._update_lock = threading.Lock()
        self._github_update = None
        self._design_files: Dict[str, Dict[str, Any]] = {}
        self._design_lock = threading.Lock()

    def bind(self, window: webview.Window) -> None:
        self.window = window

    def _open_dialog(self, multiple: bool = False, file_types=SEQUENCE_TYPES) -> List[str]:
        if not self.window:
            return []
        result = self.window.create_file_dialog(
            webview.OPEN_DIALOG,
            allow_multiple=multiple,
            file_types=file_types,
        )
        return list(result or [])

    def _save_dialog(self, filename: str, file_types: Sequence[str]) -> Optional[str]:
        if not self.window:
            return None
        result = self.window.create_file_dialog(
            webview.SAVE_DIALOG,
            save_filename=filename,
            file_types=tuple(file_types),
        )
        if not result:
            return None
        if isinstance(result, (list, tuple)):
            return str(result[0]) if result else None
        return str(result)

    def choose_parent(self) -> Optional[Dict[str, Any]]:
        paths = self._open_dialog(False)
        return self._file_info(paths[0]) if paths else None

    def choose_backbone(self) -> Optional[Dict[str, Any]]:
        """Choose a per-insert backbone override from the spreadsheet cell."""
        paths = self._open_dialog(False)
        return self._file_info(paths[0]) if paths else None

    def backbone_from_path(self, path: str) -> Dict[str, Any]:
        if not path or not Path(path).is_file():
            raise ValueError("Drop a local FASTA or SnapGene .dna file, or click the cell to choose one")
        return self._file_info(path)

    def register_drop_targets(self, wps_target_ids: List[str]) -> bool:
        """Attach native macOS drop handlers so Finder paths reach WebKit cells."""
        if not self.window:
            return False
        for index, element in enumerate(self.window.dom.get_elements(".backbone-cell")):
            element.on(
                "drop",
                DOMEventHandler(
                    lambda event, row_index=index: self._receive_dropped_backbone(row_index, event),
                    prevent_default=True,
                ),
            )
        for target_id, element in zip(wps_target_ids, self.window.dom.get_elements(".wps-cell")):
            element.on(
                "drop",
                DOMEventHandler(
                    lambda event, target=target_id: self._receive_dropped_wps(target, event),
                    prevent_default=True,
                ),
            )
        return True

    def _receive_dropped_backbone(self, row_index: int, event: Dict[str, Any]) -> None:
        try:
            files = event.get("dataTransfer", {}).get("files", [])
            path = files[0].get("pywebviewFullPath", "") if files else ""
            if not path:
                raise ValueError("No local sequence file was received from Finder")
            backbone = self.backbone_from_path(path)
            if self.window:
                self.window.evaluate_js(
                    f"window.nativeBackboneDropped({row_index}, {json.dumps(backbone)})"
                )
        except Exception as exc:
            if self.window:
                self.window.evaluate_js(
                    f"window.nativeBackboneDropFailed({row_index}, {json.dumps(str(exc))})"
                )

    def _receive_dropped_wps(self, target_id: str, event: Dict[str, Any]) -> None:
        try:
            files = event.get("dataTransfer", {}).get("files", [])
            path = files[0].get("pywebviewFullPath", "") if files else ""
            if not path:
                raise ValueError("No local sequence file was received from Finder")
            wps = self.backbone_from_path(path)
            if self.window:
                self.window.evaluate_js(
                    f"window.nativeWpsDropped({json.dumps(target_id)}, {json.dumps(wps)})"
                )
        except Exception as exc:
            if self.window:
                self.window.evaluate_js(
                    f"window.nativeWpsDropFailed({json.dumps(target_id)}, {json.dumps(str(exc))})"
                )

    def choose_inserts(self) -> List[Dict[str, Any]]:
        rows: List[Dict[str, Any]] = []
        for index, path in enumerate(self._open_dialog(True, INSERT_TYPES), 1):
            if Path(path).suffix.lower() in {".csv", ".tsv", ".txt"}:
                rows.extend(_table_insert_rows(Path(path).read_text(encoding="utf-8-sig"), Path(path).name))
            elif Path(path).suffix.lower() == ".xlsx":
                rows.extend(_xlsx_insert_rows(path))
            else:
                rows.append(self._file_info(path, index))
        return rows

    def parse_pasted_inserts(self, text: str) -> List[Dict[str, Any]]:
        return _table_insert_rows(text, "Pasted table")

    def choose_consensuses(self) -> List[Dict[str, Any]]:
        return [self._file_info(path) for path in self._open_dialog(True)]

    def choose_wps(self) -> Optional[Dict[str, Any]]:
        paths = self._open_dialog(False)
        return self._file_info(paths[0]) if paths else None

    def choose_design_plan(self) -> Optional[Dict[str, Any]]:
        paths = self._open_dialog(False, ("Golden Gate design plans (*.yaml;*.yml)", "All files (*.*)"))
        if not paths:
            return None
        path = Path(paths[0])
        return {"path": str(path), "filename": path.name}

    def choose_annotated_design_map(self) -> Optional[Dict[str, Any]]:
        paths = self._open_dialog(False, DESIGN_MAP_TYPES)
        if not paths:
            return None
        return self.annotated_design_map_from_path(paths[0])

    def annotated_design_map_from_path(self, path: str) -> Dict[str, Any]:
        path = Path(path).expanduser().resolve()
        if path.suffix.lower() != ".dna" or not path.is_file():
            raise ValueError("Choose or drop a local SnapGene .dna file")
        plan = plan_from_annotated_snapgene(path)
        fragments = plan["fragments"]
        return {
            "path": str(path), "filename": path.name,
            "fragmentCount": len(fragments),
            "fixedCount": sum(item.get("annotation_kind") == "fixed" for item in fragments),
            "variableCount": sum(item.get("annotation_kind") == "variable" for item in fragments),
            "variableFragments": [{"id": item["id"], "name": item["name"]} for item in fragments if item.get("annotation_kind") == "variable"],
            "backboneName": next(item["name"] for item in fragments if item.get("is_backbone")),
            "junctionCount": len(plan["junction_candidates"]),
        }

    def register_design_drop_target(self) -> bool:
        if not self.window:
            return False
        elements = self.window.dom.get_elements("#pick-design-map")
        for element in elements:
            if getattr(self, "_design_drop_element", None) == element:
                continue
            element.on("drop", DOMEventHandler(self._receive_dropped_design, prevent_default=True))
            self._design_drop_element = element
        return bool(elements)

    def _receive_dropped_design(self, event: Dict[str, Any]) -> None:
        try:
            files = event.get("dataTransfer", {}).get("files", [])
            if len(files) != 1:
                raise ValueError("Drop one annotated SnapGene .dna file at a time")
            path = files[0].get("pywebviewFullPath", "")
            if not path:
                raise ValueError("No local file path was received. Click to choose the file.")
            result = self.annotated_design_map_from_path(path)
            if self.window:
                self.window.evaluate_js(f"window.nativeDesignDropped({json.dumps(result)})")
        except Exception as exc:
            if self.window:
                self.window.evaluate_js(f"window.nativeDesignDropFailed({json.dumps(str(exc))})")

    def run_golden_gate_design(self, plan_path: str, output_folder: Optional[str] = None) -> Dict[str, Any]:
        try:
            if not plan_path:
                raise ValueError("Choose a design plan first")
            return self._temporary_design(lambda root: run_plan(plan_path, root))
        except Exception as exc:
            return {"ok": False, "error": str(exc), "detail": traceback.format_exc(limit=5)}

    def run_annotated_golden_gate_design(self, map_path: str, output_folder: Optional[str] = None, enzyme: str = "BsmBI", variable_texts: Optional[Dict[str, str]] = None, project_state: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        try:
            if not map_path:
                raise ValueError("Choose an annotated SnapGene map first")
            # Legacy output folders are ignored: save inside the app library;
            # exporting ordinary files still requires an explicit download.
            state = copy.deepcopy(project_state or {})
            original = (state.get("goldenGateDesign") or {}).get("map") or {}
            encoded = original.get("embeddedSnapGene") if original.get("path") == map_path else None
            data = base64.b64decode(encoded, validate=True) if encoded else Path(map_path).read_bytes()
            # Generate from the exact embedded input snapshot, not a mutable
            # external file. Reopened projects can be rerun without the source.
            with tempfile.TemporaryDirectory(prefix="cloning-project-input-") as folder:
                source = Path(folder) / Path(map_path).name
                source.write_bytes(data)
                info = self.annotated_design_map_from_path(str(source))
                info.update(path=map_path, embeddedSnapGene=base64.b64encode(data).decode("ascii"))
                state["goldenGateDesign"] = {"map": info, "enzyme": enzyme, "variableTexts": variable_texts or {}}
                state["projectName"] = _project_name(state.get("projectName"))
                state["projectId"] = state.get("projectId") or uuid.uuid4().hex
                def generate(root):
                    result = run_annotated_snapgene_design(source, root, enzyme, variable_texts)
                    result.update(projectState=state, project=state["projectName"])
                    return result
                return self._temporary_design(generate)
        except Exception as exc:
            return {"ok": False, "error": str(exc), "detail": traceback.format_exc(limit=5)}

    def _temporary_design(self, generate) -> Dict[str, Any]:
        """Clean renderer staging, then commit a complete run to the app library."""
        with self._design_lock:
            with tempfile.TemporaryDirectory(prefix="cloning-companion-design-") as workspace:
                result = _design_previews(generate(Path(workspace)))
                output = Path(result["outputDir"])
                def relative(value):
                    if isinstance(value, dict):
                        return {k: relative(v) for k, v in value.items()}
                    if isinstance(value, list):
                        return [relative(v) for v in value]
                    if isinstance(value, str) and value.startswith(str(output) + "/"):
                        return str(Path(value).relative_to(output))
                    return value
                files, by_path, downloads = {}, {}, []
                for path in sorted(output.rglob("*")):
                    if not path.is_file():
                        continue
                    name = str(path.relative_to(output))
                    key = uuid.uuid4().hex
                    data = path.read_bytes()
                    if path.name == "design.json":
                        report = relative(json.loads(data))
                        report.pop("outputDir", None)
                        if result.get("projectState"):
                            original_path = result["projectState"]["goldenGateDesign"]["map"]["path"]
                            report.update(scaffold=original_path, plan=original_path)
                        data = json.dumps(report, indent=2).encode("utf-8")
                    group = ("plasmids" if name.startswith("plasmids/") else
                             "orders" if path.name in {"synthesis_order.tsv", "synthesis_order.fasta", "pcr_primers.tsv", "assembly_recipe.tsv"} else
                             "graphics" if path.suffix.lower() in {".png", ".svg"} else "other")
                    info = {"id": key, "filename": path.name, "relativePath": name, "group": group, "bytes": len(data)}
                    files[key] = {**info, "data": data}
                    downloads.append(info)
                    by_path[str(path)] = key
                for plasmid in result.get("plasmids", []):
                    plasmid["downloadId"] = by_path[plasmid["path"]]
                for graphic in result.get("graphics", []):
                    graphic["downloadId"] = by_path[graphic["sourcePath"]]
                result = relative(result)
                result.pop("outputDir", None)
                result.update(downloads=downloads, temporary=True)
            result.update(savedRun=design_library.save(application_support(), result, files), temporary=False)
            # Replace the active view only after a successful durable save.
            self._design_files = files
            return result

    def list_saved_designs(self) -> Dict[str, Any]:
        try:
            return {"ok": True, "runs": design_library.list_runs(application_support())}
        except Exception as exc:
            return {"ok": False, "error": f"Could not read the design library: {exc}"}

    def open_saved_design(self, run_id: str) -> Dict[str, Any]:
        try:
            with self._design_lock:
                result, files = design_library.load(application_support(), run_id)
                self._design_files = files
                if result.get("projectState"):
                    self.current_project = None
                return result
        except Exception as exc:
            return {"ok": False, "error": f"Could not open the saved design: {exc}"}

    def download_design_file(self, download_id: str) -> Dict[str, Any]:
        item = self._design_files.get(download_id)
        if item is None:
            return {"ok": False, "error": "This project is not open. Reopen it from Saved projects, then download again."}
        suffix = Path(item["filename"]).suffix
        path = self._save_dialog(item["filename"], (f"Output files (*{suffix})", "All files (*.*)"))
        if not path:
            return {"ok": False, "cancelled": True}
        try:
            Path(path).write_bytes(item["data"])
            return {"ok": True, "path": path}
        except OSError as exc:
            return {"ok": False, "error": str(exc)}

    def download_design_bundle(self, group: str = "all") -> Dict[str, Any]:
        if group not in {"all", "plasmids", "orders", "graphics"}:
            return {"ok": False, "error": "Unknown output group"}
        items = [item for item in self._design_files.values() if group == "all" or item["group"] == group]
        if not items:
            return {"ok": False, "error": "Generate a design before downloading outputs."}
        path = self._save_dialog(f"cloning-design-{group}.zip", ("ZIP archives (*.zip)",))
        if not path:
            return {"ok": False, "cancelled": True}
        try:
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                for item in items:
                    archive.writestr(item["relativePath"], item["data"])
            Path(path).write_bytes(buffer.getvalue())
            return {"ok": True, "path": path}
        except OSError as exc:
            return {"ok": False, "error": str(exc)}

    def _file_info(self, path: str, order: Optional[int] = None) -> Dict[str, Any]:
        name, sequence = _first_sequence(path)
        result = {
            "path": path,
            "filename": Path(path).name,
            "recordName": name,
            "length": len(sequence),
            "targetId": infer_target_id(Path(path).name),
        }
        if order is not None:
            result.update({"order": order, "noncoding5": 0, "noncoding3": 0})
        return result

    def run_verification(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        try:
            parent_path = str(payload.get("parentPath", ""))
            parent = _first_sequence(parent_path)[1] if parent_path else ""
            insert_rows = payload.get("inserts", [])
            consensus_rows = payload.get("consensuses", [])
            results: Dict[str, Tuple[AssemblyResult, VerificationResult]] = {}
            failures: List[Dict[str, str]] = []
            seen: set = set()
            for consensus_row in consensus_rows:
                target_id = str(consensus_row.get("targetId", "")).strip()
                if not target_id:
                    failures.append({"targetId": consensus_row.get("filename", "Consensus"), "error": "Target ID is empty"})
                    continue
                if target_id in seen:
                    failures.append({"targetId": target_id, "error": "More than one WPS consensus uses this target ID"})
                    continue
                seen.add(target_id)
                try:
                    matching = [row for row in insert_rows if str(row.get("targetId", "")).strip() == target_id]
                    if not matching:
                        raise AssemblyError("No insert sequence files are paired with this target")
                    matching.sort(key=lambda row: int(row.get("order", 0)))
                    assigned_backbones = [
                        row.get("backbone", {}) for row in matching
                        if isinstance(row.get("backbone"), dict) and row.get("backbone", {}).get("path")
                    ]
                    unique_backbones = {str(item["path"]): item for item in assigned_backbones}
                    if len(unique_backbones) > 1:
                        raise AssemblyError("Inserts for the same target have different assigned backbones")
                    target_parent = (
                        _first_sequence(next(iter(unique_backbones)))[1]
                        if unique_backbones else parent
                    )
                    if not target_parent:
                        raise AssemblyError("Assign a backbone in the insert spreadsheet")
                    inserts: List[Tuple[str, str]] = []
                    noncoding: Dict[str, Tuple[int, int]] = {}
                    for row in matching:
                        sequence = row.get("sequence") or _first_sequence(row["path"])[1]
                        name = row["filename"]
                        inserts.append((name, sequence))
                        noncoding[name] = (
                            max(0, int(row.get("noncoding5", 0))),
                            max(0, int(row.get("noncoding3", 0))),
                        )
                    if payload["method"] == "Gibson":
                        assembly = assemble_gibson(
                            target_parent,
                            inserts,
                            target_id,
                            max(10, int(payload.get("armLength", 15))),
                            noncoding,
                        )
                    else:
                        assembly = assemble_golden_gate(
                            target_parent,
                            inserts,
                            target_id,
                            payload.get("enzyme", "BsaI"),
                            noncoding,
                        )
                    _, observed = _first_sequence(consensus_row["path"])
                    verification = verify_consensus(assembly, observed)
                    results[target_id] = (assembly, verification)
                except Exception as exc:
                    failures.append({"targetId": target_id, "error": str(exc)})
            self.latest_results = results
            return {
                "ok": True,
                "results": [_serialize_result(*results[target_id]) for target_id in results],
                "failures": failures,
            }
        except Exception as exc:
            return {"ok": False, "error": str(exc), "detail": traceback.format_exc(limit=5)}

    def save_project(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        return self._save_project(payload, False)

    def clear_verification_inputs(self) -> bool:
        """Discard verification exports while keeping the current project and design."""
        self.latest_results = {}
        return True

    def start_new_project(self) -> bool:
        self.current_project = None
        self._design_files = {}
        self.latest_results = {}
        return True

    def _save_project(self, payload: Dict[str, Any], as_new: bool) -> Dict[str, Any]:
        name = _project_name(payload.get("projectName"))
        path = None if as_new else self.current_project
        if not path:
            path = self._save_dialog(f"{_safe_filename(name)}.{PROJECT_EXTENSION}", PROJECT_TYPES)
        if not path:
            return {"ok": False, "cancelled": True}
        if not path.endswith(f".{PROJECT_EXTENSION}"):
            path += f".{PROJECT_EXTENSION}"
        try:
            state = copy.deepcopy(payload)
            state["projectName"] = name
            design = state.get("goldenGateDesign") or {}
            source = design.get("map") or {}
            if source.get("path") and not source.get("embeddedSnapGene"):
                source["embeddedSnapGene"] = base64.b64encode(Path(source["path"]).read_bytes()).decode("ascii")
            document = {"format": "CloningCompanion Project", "version": 2,
                        "savedAt": datetime.now(timezone.utc).isoformat(), "state": state}
            if design.get("savedRunId"):
                result, files = design_library.load(application_support(), design["savedRunId"])
                if not _same_design_inputs(state, result.get("projectState") or {}):
                    raise ValueError("Inputs have changed since these outputs were generated. Rerun the design before saving outputs with this project.")
                document["designOutputs"] = design_library.pack({**result, "projectState": state, "project": state.get("projectName") or result["project"]}, files)
            # Atomic replacement leaves an existing project intact on failure.
            with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=Path(path).parent, prefix=".cloning-project-", delete=False) as handle:
                staging = Path(handle.name)
                try:
                    json.dump(document, handle, indent=2)
                except Exception:
                    staging.unlink(missing_ok=True)
                    raise
            try:
                staging.replace(path)
            finally:
                staging.unlink(missing_ok=True)
        except Exception as exc:
            return {"ok": False, "error": f"Could not save project: {exc}"}
        self.current_project = path
        self._remember_project(path)
        return {"ok": True, "path": path, "projectName": name, "recent": self.recent_projects()}

    def save_project_as(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        return self._save_project(payload, True)

    def open_project(self, path: Optional[str] = None) -> Dict[str, Any]:
        if not path:
            paths = self._open_dialog(False, PROJECT_TYPES)
            if not paths:
                return {"ok": False, "cancelled": True}
            path = paths[0]
        try:
            document = json.loads(Path(path).read_text(encoding="utf-8"))
            state = document["state"]
            output = None
            if document.get("designOutputs"):
                output, files = design_library.unpack(document["designOutputs"])
                if not _same_design_inputs(state, output.get("projectState") or {}):
                    raise ValueError("Project inputs do not match the embedded outputs; regenerate them from the intended inputs.")
                output["projectState"] = copy.deepcopy(state)
                run = design_library.save(application_support(), output, files)
                output["savedRun"] = run
                state.setdefault("goldenGateDesign", {})["savedRunId"] = run["id"]
                self._design_files = files
            all_paths = [state.get("parentPath", "")]
            all_paths.extend(row.get("path", "") for row in state.get("inserts", []))
            all_paths.extend(
                row.get("backbone", {}).get("path", "")
                for row in state.get("inserts", [])
                if isinstance(row.get("backbone"), dict)
            )
            all_paths.extend(row.get("path", "") for row in state.get("consensuses", []))
            design_map = (state.get("goldenGateDesign") or {}).get("map") or {}
            if not design_map.get("embeddedSnapGene"):
                all_paths.append(design_map.get("path", ""))
            missing = [item for item in all_paths if item and not Path(item).exists()]
            self.current_project = path
            self._remember_project(path)
            return {"ok": True, "path": path, "state": state, "designResult": output, "missing": missing, "recent": self.recent_projects()}
        except Exception as exc:
            return {"ok": False, "error": f"Could not open project: {exc}"}

    def recent_projects(self) -> List[Dict[str, str]]:
        recent_path = application_support() / "recent-projects.json"
        if not recent_path.exists():
            return []
        try:
            paths = json.loads(recent_path.read_text(encoding="utf-8"))
        except Exception:
            return []
        return [
            {"path": path, "name": Path(path).stem}
            for path in paths
            if isinstance(path, str) and Path(path).exists()
        ][:8]

    def _remember_project(self, path: str) -> None:
        paths = [item["path"] for item in self.recent_projects() if item["path"] != path]
        paths.insert(0, path)
        (application_support() / "recent-projects.json").write_text(
            json.dumps(paths[:8], indent=2), encoding="utf-8"
        )

    def export_summary(self) -> Dict[str, Any]:
        if not self.latest_results:
            return {"ok": False, "error": "Run verification before exporting"}
        path = self._save_dialog("plasmid-verification-summary.csv", ("CSV files (*.csv)",))
        if not path:
            return {"ok": False, "cancelled": True}
        categories = ["backbone", "junction", "non-coding", "silent", "missense", "nonsense", "frameshift"]
        fields = ["target_id", "verdict", "method", "target_bp", "WPS_bp", "mutations", *categories]
        with Path(path).open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fields)
            writer.writeheader()
            for target_id, (assembly, result) in self.latest_results.items():
                counts = result.counts
                row = {
                    "target_id": target_id,
                    "verdict": result.verdict,
                    "method": assembly.method,
                    "target_bp": result.expected_length,
                    "WPS_bp": result.observed_length,
                    "mutations": len(result.mutations),
                }
                row.update({category: counts.get(category, 0) for category in categories})
                writer.writerow(row)
        return {"ok": True, "path": path}

    def export_mutations(self) -> Dict[str, Any]:
        if not self.latest_results:
            return {"ok": False, "error": "Run verification before exporting"}
        path = self._save_dialog("plasmid-verification-mutations.csv", ("CSV files (*.csv)",))
        if not path:
            return {"ok": False, "cancelled": True}
        fields = [
            "target_id", "verdict", "position", "mutation", "category", "region",
            "source", "target", "observed", "protein_change", "context",
        ]
        with Path(path).open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fields)
            writer.writeheader()
            for target_id, (_, result) in self.latest_results.items():
                for mutation in result.mutations:
                    row = {"target_id": target_id, "verdict": result.verdict, **mutation.as_dict()}
                    writer.writerow(row)
        return {"ok": True, "path": path}

    def export_targets(self) -> Dict[str, Any]:
        if not self.latest_results:
            return {"ok": False, "error": "Run verification before exporting"}
        path = self._save_dialog("generated-targets.zip", ("ZIP archives (*.zip)",))
        if not path:
            return {"ok": False, "cancelled": True}
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for target_id, (assembly, _) in self.latest_results.items():
                archive.writestr(
                    f"{_safe_filename(target_id)}_target.fasta",
                    to_fasta(target_id, assembly.sequence),
                )
        return {"ok": True, "path": path}

    def app_info(self) -> Dict[str, Any]:
        info = {
            "name": APP_NAME,
            "version": APP_VERSION,
            "python": platform.python_version(),
            "enzymes": list(TYPE_IIS_ENZYMES),
            "recent": self.recent_projects(),
        }
        library = self.list_saved_designs()
        info["savedDesigns"] = library.get("runs", [])
        if not library["ok"]:
            info["designLibraryError"] = library["error"]
        pending = local_updates.update_directory() / "pending-session.json"
        if pending.exists():
            try:
                job = Path(json.loads(pending.read_text())["job"])
                job.resolve().relative_to(local_updates.update_directory().resolve())
                session = json.loads((job.parent / "session.json").read_text())
                self.current_project = session.get("path")
                info["updateRecovery"] = {"ok": True, "state": session["state"], "path": self.current_project}
                run_id = (session["state"].get("goldenGateDesign") or {}).get("savedRunId")
                if run_id:
                    result, files = design_library.load(application_support(), run_id)
                    self._design_files = files
                    info["updateRecovery"]["designResult"] = result
            except Exception as exc:
                info["updateRecoveryError"] = f"Could not restore update session: {exc}"
        return info

    def acknowledge_update_recovery(self) -> None:
        pending = local_updates.update_directory() / "pending-session.json"
        if pending.exists():
            pending.replace(pending.with_name("last-restored-session.json"))

    def check_for_updates(self) -> Dict[str, Any]:
        try:
            self._github_update = None
            target = local_updates.running_bundle()
            local = local_updates.check_update(APP_VERSION, local_updates.update_directory(), target)
            if local.get("available"):
                local["channel"] = "Local build"
                return local
            result, envelope = github_updates.check_update(APP_VERSION, target)
            if result.get("available") and result.get("canInstall"):
                self._github_update = envelope
            return result
        except Exception as exc:
            return {"ok": False, "error": f"Could not check updates: {exc}. Your installed app is unchanged; try again later."}

    def install_update(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        if not self._update_lock.acquire(blocking=False):
            return {"ok": False, "error": "An update is already in progress"}
        try:
            target = local_updates.running_bundle()
            if target is None or not self.window:
                raise ValueError("Updates can only be installed from the desktop app")
            directory = local_updates.update_directory()
            manifest = github_updates.download_update(self._github_update, directory, APP_VERSION) if self._github_update else None
            job = local_updates.stage_update(target, directory, APP_VERSION,
                                             {"state": payload, "path": self.current_project}, manifest=manifest)
            local_updates.launch_helper(job)
            # The helper is ready and waits for this process to exit before
            # touching the installed app. Keep the lock until termination.
            threading.Timer(1.0, self.window.destroy).start()
            return {"ok": True, "recovery": str(job.parent / "recovery.plasmidverify")}
        except Exception as exc:
            self._update_lock.release()
            return {"ok": False, "error": f"Update not installed: {exc}"}

    def reveal_in_finder(self, path: str) -> bool:
        if not path or not Path(path).exists():
            return False
        subprocess.run(["open", "-R", path], check=False)
        return True


api = NativeAPI()


def _js_call(function: str) -> None:
    if api.window:
        api.window.evaluate_js(function)


def menu_open_project() -> None:
    result = api.open_project()
    if result.get("ok"):
        _js_call(f"window.loadProjectFromNative({json.dumps(result)})")


def menu_save_project() -> None:
    _js_call("window.saveProjectFromMenu()")


def menu_export_summary() -> None:
    result = api.export_summary()
    _js_call(f"window.nativeExportFinished({json.dumps(result)})")


def menu_export_mutations() -> None:
    result = api.export_mutations()
    _js_call(f"window.nativeExportFinished({json.dumps(result)})")


def menu_export_targets() -> None:
    result = api.export_targets()
    _js_call(f"window.nativeExportFinished({json.dumps(result)})")


def menu_about() -> None:
    _js_call("window.showModal('about')")


def menu_help() -> None:
    _js_call("window.showModal('help')")


def menu_updates() -> None:
    _js_call("window.checkForUpdates()")


def menu_new() -> None:
    _js_call("window.newProject()")


def main() -> None:
    index = resource_path("mac_app/index.html")
    window = webview.create_window(
        APP_NAME,
        index.as_uri(),
        js_api=api,
        width=1280,
        height=820,
        min_size=(980, 680),
        background_color="#F8FAF6",
        confirm_close=False,
    )
    api.bind(window)
    menu = [
        Menu(
            "File",
            [
                MenuAction("New Project", menu_new),
                MenuAction("Open Project…", menu_open_project),
                MenuAction("Save Project…", menu_save_project),
                MenuSeparator(),
                MenuAction("Export Summary CSV…", menu_export_summary),
                MenuAction("Export Mutations CSV…", menu_export_mutations),
                MenuAction("Export Target FASTAs…", menu_export_targets),
            ],
        ),
        Menu(
            "Help",
            [
                MenuAction("CloningCompanion Help", menu_help),
                MenuAction("Check for Updates…", menu_updates),
                MenuSeparator(),
                MenuAction("About CloningCompanion", menu_about),
            ],
        ),
    ]
    webview.start(menu=menu, debug=False, private_mode=False)


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--apply-local-update":
        sys.exit(local_updates.apply_update_job(Path(sys.argv[2])))
    main()
